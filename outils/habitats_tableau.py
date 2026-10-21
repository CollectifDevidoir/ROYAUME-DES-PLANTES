# Classeur de relecture des milieux et habitats : donnees/milieux/Associations_milieux_habitats.xlsx
# Feuilles : Associations (une ligne par association, libellés en clair, filtres), Par territoire, Sans association.
# Noms français tirés des listes S et X… de index.html. Usage : python3 outils/habitats_tableau.py
import json,re,collections,pathlib,sys
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
sys.path.insert(0,str(pathlib.Path(__file__).resolve().parent));import habitats_vers_appli as h
s=(h.R/'index.html').read_text(encoding='utf-8');rows=h.lire();noms=h.especes(s)
FR={a:b for a,b in re.findall(r'\["([^"]+)",\s*"([^"]+)",\s*"[^"]+"\]',s[s.index('const S=['):s.index('const X={')])}
for a,b in re.findall(r'([A-Z][^|;"\n]+)\|([^;"\n]+)',s[s.index('const X={'):s.index('for(const t in X4)')]):FR.setdefault(a.strip(),b.strip())
CTX={'habitat_naturel':'naturel','presence_spontanee_secondaire':'spontané (milieu secondaire)','cultive_ou_plante':'cultivé ou planté'}
CONF={'verified':'vérifiée','needs_review':'à vérifier (non affichée)'}
wb=Workbook();ws=wb.active;ws.title='Associations'
H=['Nom latin','Nom français','Statut','Territoire','Sous-milieu','Code','Contexte','Force','Confiance','Extrait de la source','Source','Remarque']
ws.append(H)
for r in sorted(rows,key=lambda r:(r['Nom latin'],r['Confiance']!='verified')):
    sm=h.SM[r['Sous-milieu']]
    ws.append([r['Nom latin'],FR.get(r['Nom latin'],''),r['Statut'],h.TE[r['Territoire']]['nom'],sm['libelle'],r['Sous-milieu'],CTX[r['Contexte']],r['Force'],CONF[r['Confiance']],r['Extrait'],r['Source'],r['Remarque']])
    if r['Source']:ws.cell(ws.max_row,11).hyperlink=r['Source']
for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='456B45')
for col,w in zip('ABCDEFGHIJKL',[26,26,18,22,34,8,22,10,14,70,40,40]):ws.column_dimensions[col].width=w
for row in ws.iter_rows(min_row=2):
    for c in row:c.alignment=Alignment(wrap_text=c.column in (10,12),vertical='top')
ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
t=wb.create_sheet('Par territoire');t.append(['Territoire','Sous-milieu','Code','Espèces (associations vérifiées)'])
n=collections.Counter((r['Territoire'],r['Sous-milieu']) for r in rows if r['Confiance']=='verified')
for (te,sm),k in sorted(n.items()):t.append([h.TE[te]['nom'],h.SM[sm]['libelle'],sm,k])
for c in t[1]:c.font=Font(bold=True)
for col,w in zip('ABCD',[26,40,8,14]):t.column_dimensions[col].width=w
z=wb.create_sheet('Sans association');z.append(['Nom latin','Nom français','État'])
ok={r['Nom latin'] for r in rows if r['Confiance']=='verified'};pr={r['Nom latin'] for r in rows}
for x in sorted(noms-ok):z.append([x,FR.get(x,''),'à vérifier seulement' if x in pr else 'aucune source trouvée'])
for c in z[1]:c.font=Font(bold=True)
z.column_dimensions['A'].width=30;z.column_dimensions['B'].width=30;z.column_dimensions['C'].width=24
out=h.D/'milieux'/'Associations_milieux_habitats.xlsx';wb.save(out)
print(out.relative_to(h.R),':',len(rows),'associations,',len(ok),'espèces avec association vérifiée sur',len(noms),';',len(noms-ok),'sans')
